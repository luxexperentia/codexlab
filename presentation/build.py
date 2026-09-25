import sys, json, html
from pathlib import Path
sys.path.insert(0, '/tmp/presentation-deps')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).parent
SOURCES = {
 'A': ('Anthropic · multi-agent research · 13 Jun 2025', 'https://www.anthropic.com/engineering/multi-agent-research-system'),
 'B': ('Anthropic · compiler experiment · 5 Feb 2026', 'https://www.anthropic.com/engineering/building-c-compiler'),
 'C': ('Klarna · 2025 Form 20-F · company estimates', 'https://s205.q4cdn.com/644747736/files/doc_financials/2025/q4/Klarna-Group-plc-20-F-2025.pdf'),
 'D': ('Salesforce · customer results · 23 Jun 2025', 'https://investor.salesforce.com/news/news-details/2025/Salesforce-Launches-Agentforce-3-to-Solve-the-Biggest-Blockers-to-Scaling-AI-Agents-Visibility-and-Control/'),
 'E': ('Brynjolfsson, Li & Raymond · QJE · 2025', 'https://doi.org/10.1093/qje/qjae044'),
 'F': ('Klarna · IPO prospectus · Sep 2025', 'https://www.sec.gov/Archives/edgar/data/2003292/000200329225000052/klarnagroupplc424b4.htm'),
}

slides = [
dict(tag='EXECUTIVE BRIEFING / SEPTEMBER 2026', title='Scaling a business\nwith teams of AI agents', subtitle='The evidence, the operating model, and the decision to make now.',
 cards=[('CAPACITY','Complete more work without a matching increase in staffing.'),('COORDINATION','Give specialist agents distinct responsibilities and shared goals.'),('ECONOMICS','Scale the workflows that prove their value.')],
 takeaway='A 15-minute case for building an agent workforce.',
 notes='Open with the business constraint: demand grows faster than the organization can recruit, train and coordinate people. This presentation argues for starting a measured deployment now. It does not claim that every business needs a large agent fleet or that agents outperform people at every task. The recommendation is to establish a working capability and expand it on evidence.'),
dict(tag='01 / THE BUSINESS CONSTRAINT',title='Growth creates more work\nbefore it creates more capacity.',
 cards=[('More customers','More questions, onboarding checks, transactions and exceptions.'),('More products','More documentation, research, software changes and support.'),('More markets','More languages, operating hours and local requirements.')],
 takeaway='When every additional unit of work requires human attention, growth carries a staffing burden.',
 notes='Ask the audience to identify the queue that currently limits growth. It might be response times, onboarding, account research or internal engineering requests. The economic opportunity comes from changing the amount of human attention required per completed unit of work. This is an analytical framing, not an empirical claim about every company.'),
dict(tag='02 / WHAT THE TECHNOLOGY CHANGES',title='An agent can carry work\nthrough several steps.',
 cards=[('An assistant','Responds to a request: summarize this account or draft this reply.'),('An agent','Uses permitted tools to inspect the account, choose a next step and complete a bounded task.'),('A team of agents','Divides a larger job among specialists, combines their work and escalates exceptions.')],
 takeaway='The business value appears when completed work replaces a queue of manual steps.',
 notes='These are practical definitions for this presentation. Autonomy is bounded by permissions, tools and business rules. A team can involve a coordinator and specialist workers, or several agents coordinating through shared state. Multiple separate chatbots do not automatically constitute a coordinated team.',refs=['A']),
dict(tag='03 / WHY MULTIPLE AGENTS',title='Divide the work where\nspecialization and parallelism help.',kind='flow',
 cards=[('Coordinator','Defines the objective, assigns tasks and tracks completion.'),('Specialists','Research separate sources; inspect records; draft an answer.'),('Reviewer + owner','Check evidence and rules; approve consequential actions.')],
 takeaway='Example: a proposal team can investigate the customer, solution and delivery plan concurrently.',
 notes='This proposal workflow is an illustrative design. Independent research streams can run concurrently; the final proposal still depends on their outputs. Separate roles can keep instructions focused. A reviewer agent can share the same blind spots as the writer, so agreement is not proof of correctness. Business rules, authoritative data and accountable owners remain necessary.'),
dict(tag='04 / DIRECT MULTI-AGENT EVIDENCE',title='A research team outperformed\na single research agent.',metric='+90.2%',metric_label='performance on an internal research evaluation',
 cards=[('The comparison','Anthropic compared an Opus 4 lead with Sonnet 4 workers against a single Opus 4 agent.'),('Why it matters','Independent research streams expand coverage before synthesis.'),('The boundary','A vendor evaluation against one agent; additional compute was used. This is not a human comparison.')],
 takeaway='For suitable tasks, coordination can deliver a measurable quality gain.',refs=['A'],
 notes='This result supports a specific architecture on a specific evaluation. It does not isolate coordination from the effects of greater compute or model mix. Ask whether your workflow has independent lines of inquiry whose findings can be checked and combined.'),
dict(tag='05 / DIRECT MULTI-AGENT DEMONSTRATION',title='Sixteen agents built\na substantial compiler prototype.',metric='16 agents',metric_label='roughly two weeks · nearly 2,000 sessions · about $20,000 in API costs',
 cards=[('The output','A 100,000-line Rust-based C compiler that could build Linux 6.9 on three architectures.'),('Human contribution','A researcher designed the harness, tests and interventions that made parallel work possible.'),('The boundary','The prototype needed external tools and had quality and efficiency limitations. API cost excludes human effort.')],
 takeaway='A small supervising team can attempt a much larger body of implementation work.',refs=['B'],
 notes='Anthropic published this experiment in February 2026. The build used external assembler/linker tooling, and x86 bootstrapping still required GCC support. The artifact was not a production compiler replacement. Treat it as evidence of substantial coordinated execution, not proof that a software department can be removed.'),
dict(tag='06 / PRODUCTION ECONOMICS',title='Klarna reports substantial\ncustomer-service labor substitution.',metric='$39 million',metric_label='reported cost savings in 2024',
 cards=[('Scale','80% of customer-service chats handled by its AI assistant during 2025.'),('Labor equivalent','Over 700 full-time agents’ workload, estimated from reductions in human-handled conversations.'),('Interpretation','Routine work moved to software. The workload equivalent is not a verified count of people dismissed.')],
 takeaway='This is deployment evidence for AI automation; the filing does not establish a multi-agent architecture.',refs=['C'],
 notes='Klarna’s annual filing gives different time periods for savings and chat share; keep those dates attached to each number. These are company-reported metrics, not an independent causal evaluation. Its prospectus also describes combining scalable AI service with high-quality human support. Savings do not imply that every interaction should be automated.',refs_extra=['F']),
dict(tag='07 / GROWTH WITHOUT PROPORTIONAL STAFFING',title='At peak tax season, software\nabsorbed routine support demand.',metric='70%',metric_label='of administrative chat engagements resolved autonomously',
 cards=[('The business','1-800Accountant, a virtual accounting firm serving small businesses.'),('The setting','Critical tax weeks in 2025, when administrative demand spikes.'),('The consequence','Staff can concentrate on cases that require professional attention.')],
 takeaway='The reported result covers administrative chats; it does not mean autonomous accounting or tax advice.',refs=['D'],
 notes='Salesforce published this customer result and included confirmation from the customer’s CTO. It is a vendor-reported production example without a controlled human comparison. The case demonstrates bounded agent automation; it does not disclose evidence that coordination among multiple agents caused the result.'),
dict(tag='08 / THE HUMAN COMPARISON',title='“Better” becomes useful\nwhen the task and metric are explicit.',
 cards=[('Speed · reported','Klarna: two-minute query resolution versus 12 minutes for human agents, as of September 2024.'),('Throughput · studied','A published study of 5,172 support workers found 15% more issues resolved per hour with AI assistance.'),('Quality · conditional','Klarna reports similar satisfaction. The academic result varies across workers. Neither establishes universal superiority.')],
 takeaway='Evaluate speed, quality and cost together on the actual work your business needs done.',refs=['C','E'],
 notes='The academic result concerns people using an assistant, not autonomous agents replacing people. Klarna’s figures are observational company reports and may reflect differences in case mix. The defensible conclusion is that AI can outperform previous processes on specific measures; the organization must establish that advantage on its own workload.'),
dict(tag='09 / HOW THE ADVANTAGE SCALES',title='Lower human effort per case\nchanges the capacity of the same team.',kind='math',
 cards=[('Today','1,000 cases × 12 minutes\n= 200 human hours'),('Illustrative redesign','600 automated cases\n400 exceptions × 15 minutes\n+ 20 hours of review\n= 120 human hours'),('Capacity released','80 hours per 1,000 cases\n40% less human effort\n~1.67× throughput at fixed hours')],
 takeaway='Illustration only: assumes stable case mix and quality; excludes software and implementation costs.',
 notes='Walk through the arithmetic. Automated cases require no direct handling in this scenario; their oversight is included in the 20-hour review allowance. Exceptions take longer than the original average. At a stable 0.12 human hours per case, a 200-hour budget supports about 1,667 cases. Capacity can support growth or redeployment; it becomes cash savings only if spending actually falls. These inputs are hypothetical and are not derived from the preceding case studies.'),
dict(tag='10 / THE BUSINESS CASE',title='Count the cost\nof a successfully completed case.',
 cards=[('Include every cost','Model usage, software, integration, supervision, retries and rework.'),('Use a quality threshold','Completion must meet your accuracy, service and escalation standards.'),('Connect to value','Track avoided spending or additional profitable volume; separate both from hours released.')],
 takeaway='Unit cost = total workflow cost ÷ cases completed to the required standard.',
 notes='A cheap response that causes a second contact is not necessarily a cheap outcome. A production decision should include ongoing cost plus an explicit treatment of implementation cost, such as amortizing it over expected volume. Compare the same case types over the same period. This is a proposed evaluation framework, not a return-on-investment forecast.'),
dict(tag='11 / WHEN A TEAM EARNS ITS COST',title='Add agents only when\nthe division of work improves results.',
 cards=[('Good candidates','Independent research streams, distinct tools, large information sets and checkable outputs.'),('Poor candidates','Simple lookups, tightly dependent steps or tasks without a reliable success measure.'),('Economic reality','Anthropic reports about 15× chat token usage for its multi-agent systems. More capacity has a cost.')],
 takeaway='Compare the team against a single agent and a simpler automated workflow.',refs=['A'],
 notes='The reported token multiplier uses chat interactions as its baseline; it is not a universal price ratio or a comparison with human labor. More agents can duplicate work and introduce handoff failures. The decision criterion is the improvement in accepted outcomes per dollar on the target task.'),
dict(tag='12 / WHAT STARTING EARLIER BUYS',title='The advantage accumulates\nthrough operating experience.',kind='flow',
 cards=[('First deployment','Connect the data, define permissions and discover real exceptions.'),('Repeated execution','Build evaluation examples, measure failures and refine escalation.'),('Next workflow','Reuse integrations, operating knowledge and ownership structures.')],
 takeaway='A competitor already operating these workflows may be improving them while you are still selecting tools.',
 notes='This is a strategic inference, not a quantified finding about a specific competitor. Model access can be purchased quickly; workflow knowledge and trusted deployment practices take repeated execution. The cost of waiting depends on the value and volume of your automatable work. Do not claim that every competitor has already deployed agents.'),
dict(tag='13 / A PRACTICAL FIRST TEAM',title='Choose one queue and give\neach agent a defined responsibility.',kind='flow',
 cards=[('Intake + research','Classify the request, retrieve approved records and identify missing information.'),('Execution + checking','Prepare or perform allowed actions; check the result against rules and evidence.'),('Human owner','Handle ambiguity, approve commitments and own customer outcomes.')],
 takeaway='Example starting point: administrative customer requests with clear policies and reversible actions.',
 notes='This is an illustrative team design, not a requirement for five separate model instances. Start with the minimum architecture that works, then separate roles where specialization or parallel execution proves useful. Give each role a clear output, narrow permissions, a spending limit and a handoff path. Customer commitments and sensitive exceptions should follow the business’s existing authority structure.'),
dict(tag='14 / THE FIRST 30 DAYS',title='Build enough evidence\nto make a real scaling decision.',
 cards=[('Days 1–7 · baseline','Choose one queue and accountable owner. Measure cost, quality, volume and response time.'),('Days 8–21 · compare','Test on representative cases; compare single-agent and team designs. Run with human review.'),('Days 22–30 · decide','Use a bounded live pilot if ready. Expand only when quality and economics meet agreed thresholds.')],
 takeaway='Predefine the success criteria, spending cap and stop conditions before the pilot begins.',
 notes='This is a proposed schedule, not a promise that every integration can be completed in a month. Include difficult cases in the sample. Agree how errors, repeat contacts and escalation time count. Continue preparation if data access or quality is not ready. A failed pilot that identifies the constraint can still prevent a costly rollout.'),
dict(tag='15 / THE DECISION',title='Give one agent team\na real business objective now.',
 cards=[('Assign','An executive sponsor and an operational owner.'),('Fund','One bounded workflow, with a budget and a measurable baseline.'),('Review','A dated decision to expand, revise or stop based on accepted outcomes and cost.')],
 takeaway='The path to an agent workforce starts with one team that earns the right to scale.',
 notes='Close by asking the audience to name the queue, owner and review date. The case for urgency rests on demonstrated capability and the time required to build operational experience. The evidence supports starting now on suitable work; it does not support buying an arbitrary number of agents or promising universal human replacement.'),
dict(tag='APPENDIX / EVIDENCE REGISTER',title='What the examples establish',
 cards=[('A · research system','Internal evaluation of a coordinated team versus a single agent. Architecture and compute both differ.'),('B · compiler','Vendor research demonstration with a published artifact. Human setup and technical limits remain material.'),('C + F · Klarna','Company filings on operational savings and labor equivalence, alongside a hybrid service strategy.')],
 takeaway='Primary sources are linked below and in the speaker notes.',refs=['A','B','C','F'],
 notes='Use the source documents to check scope and measurement periods. Primary means the organization or researchers reported the result; it does not mean every metric was independently audited. Selected cases are examples, not a representative success rate for agent projects.'),
dict(tag='APPENDIX / EVIDENCE REGISTER',title='Keep the comparisons distinct',
 cards=[('D · 1-800Accountant','Customer production result published by its vendor. The metric concerns administrative chats at peak demand.'),('E · support-worker study','Peer-reviewed evidence on AI-assisted people. It is not a multi-agent or autonomous-replacement experiment.'),('Illustrations + inferences','The capacity model, rollout schedule and competitive argument are analytical proposals, not reported customer outcomes.')],
 takeaway='Evidence checked 25 September 2026. Source dates and result periods are preserved.',refs=['D','E'],
 notes='The deck deliberately separates coordinated multi-agent evidence from broader evidence that agent automation changes business economics. A business can combine those capabilities, but it must measure the incremental benefit of coordination locally.')
]

BG='101C2B'; PANEL='1C2B3D'; WHITE='F3F5F7'; MUTED='B6C3D2'; ACCENT='67DDC0'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
prs.core_properties.title='Scaling a business with teams of AI agents'
prs.core_properties.subject='Executive briefing with evidence, speaker notes and sources'
prs.core_properties.author='Business strategy briefing'

def rect(sl,x,y,w,h,color):
    s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(color); s.line.fill.background(); return s
def txt(sl,x,y,w,h,text,size=20,color=WHITE,bold=False):
    box=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(5)
    return box

for i,d in enumerate(slides):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid(); sl.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    rect(sl,.55,.42,.10,.23,ACCENT)
    txt(sl,.82,.40,11,.3,d['tag'],11,ACCENT,True)
    txt(sl,.65,.98,12.05,1.47,d['title'],34,WHITE,True)
    metric='metric' in d
    y=4.05 if metric else 3.02
    if metric:
        txt(sl,.65,2.48,5,.85,d['metric'],49,ACCENT,True)
        txt(sl,.68,3.38,12,.45,d['metric_label'],17,MUTED)
    elif 'subtitle' in d:
        txt(sl,.68,2.50,12,.45,d['subtitle'],18,MUTED)
    for j,(head,body) in enumerate(d['cards']):
        x=.65+j*4.1; h=1.91 if metric else 2.76
        rect(sl,x,y,3.83,h,PANEL)
        txt(sl,x+.2,y+.19,3.43,.46,head,17,ACCENT,True)
        txt(sl,x+.2,y+.77,3.40,h-.85,body,16 if metric else 19,WHITE)
        if d.get('kind')=='flow' and j<2:
            txt(sl,x+3.87,y+.85,.23,.4,'›',23,ACCENT,True)
    txt(sl,.67,6.18,12,.54,d['takeaway'],17,WHITE,True)
    refs=d.get('refs',[])
    for k,key in enumerate(refs):
        label,url=SOURCES[key]
        b=txt(sl,.67,6.92+k*.12,11.7,.15,f'[{key}] {label}',8,MUTED)
        b.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    txt(sl,12.1,7.08,.55,.2,f'{i+1:02}',10,MUTED)
    notes=d['notes']+'\n\nSources:\n'+'\n'.join(f'[{k}] {SOURCES[k][0]}\n{SOURCES[k][1]}' for k in refs+d.get('refs_extra',[]))
    sl.notes_slide.notes_text_frame.text=notes
prs.save(OUT/'AI_Agent_Teams.pptx')

md=['# Scaling a business with teams of AI agents','', '18 slides · approximately 15 minutes plus discussion · evidence checked 25 September 2026','']
sections=[]
for i,d in enumerate(slides):
    title=d['title'].replace('\n',' ')
    md += [f'## {i+1}. {title}','',d['tag'],'']
    metric_html=''
    if 'metric' in d:
        md += [f"**{d['metric']}** — {d['metric_label']}",'']
        metric_html=f'<div class="metric">{html.escape(d["metric"])}</div><div class="metric-label">{html.escape(d["metric_label"])}</div>'
    for h,b in d['cards']: md += [f'**{h}:** {b.replace(chr(10), " · ")}','']
    md += [f'**{d["takeaway"]}**','','### Speaker notes','',d['notes'],'']
    links=[]
    for k in d.get('refs',[])+d.get('refs_extra',[]):
        name,url=SOURCES[k]; md += [f'[{k}: {name}]({url})','']; links.append(f'<a href="{url}" target="_blank" rel="noopener">[{k}] {html.escape(name)}</a>')
    cards=''.join(f'<article><h3>{html.escape(h)}</h3><p>{html.escape(b).replace(chr(10),"<br>")}</p></article>' for h,b in d['cards'])
    sections.append(f'<section class="slide {"has-metric" if metric_html else ""}" id="slide-{i+1}"><div class="tag">{html.escape(d["tag"])}</div><h1>{html.escape(d["title"]).replace(chr(10),"<br>")}</h1>{metric_html}<div class="cards">{cards}</div><p class="takeaway">{html.escape(d["takeaway"])}</p><footer>{"<br>".join(links)}<span>{i+1:02} / {len(slides)}</span></footer><aside><strong>Speaker notes</strong><p>{html.escape(d["notes"])}</p></aside></section>')
(OUT/'Speaker_Notes_and_Sources.md').write_text('\n'.join(md))
(OUT/'slides.json').write_text(json.dumps(slides,indent=2))
css='''*{box-sizing:border-box}body{margin:0;background:#09111c;color:#F3F5F7;font-family:Arial,sans-serif}.slide{max-width:1280px;min-height:720px;margin:24px auto;background:#101C2B;padding:40px 56px;position:relative;break-after:page}.tag{color:#67DDC0;font-size:13px;letter-spacing:2px}h1{font-size:43px;line-height:1.1;margin:30px 0 42px}.cards{display:grid;grid-template-columns:repeat(3,1fr);gap:24px}article{background:#1C2B3D;padding:22px;min-height:220px}h3{color:#67DDC0;margin:0 0 22px;font-size:21px}article p{font-size:23px;line-height:1.3;margin:0}.takeaway{font-size:21px;font-weight:600;margin-top:32px}footer{color:#B6C3D2;font-size:11px;margin-top:25px}footer span{float:right}a{color:inherit}.metric{font-size:65px;color:#67DDC0;font-weight:700}.metric-label{font-size:19px;color:#B6C3D2;margin:8px 0 25px}.has-metric h1{margin-bottom:18px}.has-metric article{min-height:160px}.has-metric article p{font-size:19px}.has-metric h3{font-size:19px;margin-bottom:14px}aside{display:none;margin-top:25px;border-top:1px solid #536478;padding-top:20px;line-height:1.5;color:#B6C3D2}body.notes aside{display:block}nav{position:fixed;bottom:12px;right:20px;z-index:4;display:flex;gap:8px}button{background:#67DDC0;border:0;padding:10px 16px;cursor:pointer;color:#101C2B;border-radius:4px}@media(max-width:800px){.slide{padding:30px 25px}h1{font-size:32px}.cards{grid-template-columns:1fr}article{min-height:0}article p{font-size:19px}}@media print{@page{size:landscape;margin:0}body{background:#101C2B;-webkit-print-color-adjust:exact;print-color-adjust:exact}.slide{margin:0;width:100vw;height:100vh;min-height:0;max-width:none;padding:30px 45px}nav,aside{display:none!important}h1{font-size:34px;margin:20px 0 28px}article p{font-size:18px}article{min-height:180px}.metric{font-size:50px}.has-metric article p{font-size:16px}.has-metric article{min-height:145px}.takeaway{font-size:18px;margin-top:22px}footer{font-size:9px}}'''
script='''let current=0;const slides=[...document.querySelectorAll('.slide')];function go(n){current=Math.max(0,Math.min(slides.length-1,n));slides[current].scrollIntoView({behavior:'smooth'});}document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();go(current+1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();go(current-1)}if(e.key==='n')document.body.classList.toggle('notes')});const observer=new IntersectionObserver(entries=>{for(const e of entries)if(e.isIntersecting)current=slides.indexOf(e.target)},{threshold:0.6});slides.forEach(s=>observer.observe(s));'''
page=f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Scaling with AI agent teams</title><style>{css}</style></head><body>{"".join(sections)}<nav aria-label="Presentation controls"><button onclick="go(current-1)">Previous</button><button onclick="go(current+1)">Next</button><button onclick="document.body.classList.toggle(\'notes\')">Notes</button></nav><script>{script}</script></body></html>'
(OUT/'AI_Agent_Teams.html').write_text(page)
print(f'Created {len(slides)} slides, PowerPoint, HTML and speaker notes in {OUT}')
